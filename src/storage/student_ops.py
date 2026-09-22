from __future__ import annotations

import csv
from collections import namedtuple
from pathlib import Path

from src.storage.db import db
from src.storage.models import Classroom, Student
from src.utils.persian_utils import to_ascii_digits

TEMPLATE_HEADERS = ("نام", "نام خانوادگی", "کد ملی", "کلاس")

ImportRow = namedtuple(
    "ImportRow", ("line", "first_name", "last_name", "national_id", "classroom_name")
)
RowError = namedtuple("RowError", ("line", "national_id", "reason"))
ImportResult = namedtuple("ImportResult", ("created", "skipped", "errors"))


def _rows_from_xlsx(path: Path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for raw in workbook.active.iter_rows(values_only=True):
            yield ["" if cell is None else str(cell) for cell in raw]
    finally:
        workbook.close()


def _rows_from_csv(path: Path):
    with open(path, newline="", encoding="utf-8-sig") as handle:
        yield from csv.reader(handle)


def read_roster(path) -> list[ImportRow]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        source = _rows_from_xlsx(path)
    elif suffix == ".csv":
        source = _rows_from_csv(path)
    else:
        raise ValueError("فرمت فایل پشتیبانی نمی‌شود؛ فقط xlsx و csv.")

    rows: list[ImportRow] = []
    columns: dict[str, int] | None = None
    for number, raw in enumerate(source, start=1):
        cells = [(cell or "").strip() for cell in raw]
        if columns is None:
            columns = {header: index for index, header in enumerate(cells)}
            missing = [header for header in TEMPLATE_HEADERS if header not in columns]
            if missing:
                raise ValueError("ستون «{}» پیدا نشد.".format(missing[0]))
            continue
        picked = [
            cells[columns[header]] if columns[header] < len(cells) else ""
            for header in TEMPLATE_HEADERS
        ]
        if not any(picked):
            continue
        first_name, last_name, national_id, classroom_name = picked
        rows.append(
            ImportRow(
                line=number,
                first_name=first_name,
                last_name=last_name,
                national_id=to_ascii_digits(national_id),
                classroom_name=classroom_name,
            )
        )
    if columns is None:
        raise ValueError("فایل خالی است؛ سطر عنوان پیدا نشد.")
    return rows


def write_roster(path, students) -> None:
    path = Path(path)
    suffix = path.suffix.lower()
    records = [[s.first_name, s.last_name, s.national_id, s.classroom.name] for s in students]
    if suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(TEMPLATE_HEADERS))
        for record in records:
            sheet.append(record)
        workbook.save(str(path))
    elif suffix == ".csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow(TEMPLATE_HEADERS)
            writer.writerows(records)
    else:
        raise ValueError("فرمت فایل پشتیبانی نمی‌شود؛ فقط xlsx و csv.")


def _valid_national_id(value: str) -> bool:
    return len(value) == 10 and value.isascii() and value.isdigit()


def bulk_create_students(rows, *, commit: bool = True, actor=None) -> ImportResult:
    existing_ids = {row.national_id for row in Student.select(Student.national_id)}
    rooms = {room.name: room for room in Classroom.select()}

    errors: list[RowError] = []
    skipped: list[RowError] = []
    to_create: list[tuple[ImportRow, Classroom]] = []
    seen: set[str] = set()

    for row in rows:
        if not row.first_name or not row.last_name:
            errors.append(RowError(row.line, row.national_id, "نام یا نام خانوادگی خالی است."))
            continue
        if not _valid_national_id(row.national_id):
            errors.append(RowError(row.line, row.national_id, "کد ملی باید دقیقاً ۱۰ رقم باشد."))
            continue
        if row.classroom_name not in rooms:
            errors.append(
                RowError(
                    row.line, row.national_id, "کلاس «{}» یافت نشد.".format(row.classroom_name)
                )
            )
            continue
        if row.national_id in existing_ids:
            skipped.append(
                RowError(row.line, row.national_id, "دانش‌آموزی با این کد ملی از قبل ثبت شده است.")
            )
            continue
        if row.national_id in seen:
            skipped.append(RowError(row.line, row.national_id, "کد ملی تکراری در فایل."))
            continue
        seen.add(row.national_id)
        to_create.append((row, rooms[row.classroom_name]))

    if commit and to_create:
        with db.atomic():
            for row, room in to_create:
                Student.create(
                    first_name=row.first_name,
                    last_name=row.last_name,
                    national_id=row.national_id,
                    classroom=room,
                    major=room.major,
                )
    return ImportResult(created=len(to_create), skipped=skipped, errors=errors)
